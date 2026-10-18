"""Editable PowerPoint version of slides.html, in the same Life Engine theme.

Usage (from the repository root, after make_figures.py):
    pip install python-pptx
    python Presentation/deck/make_pptx.py

Writes Presentation/Digging_Their_Own_Graves.pptx. Text uses Arial and Courier New so it renders the same on any
machine; the charts are the PNGs in figures/.
"""
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Circle, Rectangle
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "figures")
OUT = os.path.join(os.path.dirname(HERE), "Digging_Their_Own_Graves.pptx")

# the Life Engine's "neon" colours
EMPTY, PANEL, INK, MUTED, BORDER, FOOT = "0E1318", "121A22", "E6EDF3", "8B98A8", "223040", "6B7888"
SOFT = "C9D3DD"
CELL = {"mouth": "DEB14D", "producer": "15DE59", "mover": "60D4FF", "killer": "F82380", "eye": "B6C1EA",
        "claw": "FF6B35", "food": "2F7AB7"}
CLAW, PREY, KILLER, PLANT, SOIL = CELL["claw"], CELL["mover"], CELL["killer"], CELL["producer"], CELL["mouth"]
BODY_FONT, PIXEL_FONT = "Arial", "Courier New"

BODIES = {
    "plant": [(0, 0, "mouth"), (1, 1, "producer"), (-1, -1, "producer")],
    "prey": [(0, 0, "mover"), (1, 0, "eye"), (0, -1, "mouth"), (-1, 0, "mouth")],
    "clawprey": [(0, 0, "mover"), (1, 0, "eye"), (0, -1, "mouth"), (-1, 0, "mouth"), (0, 1, "claw")],
    "predator": [(0, 0, "mouth"), (0, -1, "killer"), (0, 1, "mouth"), (-1, 0, "mover"), (-2, 0, "eye"), (1, 0, "eye")],
    "eyeless": [(0, 0, "mover"), (1, 0, "killer"), (-1, 0, "claw"), (0, -1, "claw"), (0, 1, "claw")],
}


def px(v):
    """Slide coordinates are given in pixels of the 1280 x 720 HTML slides."""
    return Emu(int(v * 9525))


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


# ---------- backgrounds: the Life Engine grid, with a soil band for some slides ----------

def make_background(path, soil=False):
    fig = plt.figure(figsize=(12.8, 7.2), dpi=100)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, 1280)
    ax.set_ylim(720, 0)
    ax.axis("off")
    ax.add_patch(Rectangle((0, 0), 1280, 720, color="#" + EMPTY))
    for x in range(0, 1281, 32):
        ax.plot([x, x], [0, 720], color="#1A2430", linewidth=1)
    for y in range(0, 721, 32):
        ax.plot([0, 1280], [y, y], color="#1A2430", linewidth=1)
    # vignette, like the HTML slides
    import numpy as np
    yy, xx = np.mgrid[0:720, 0:1280]
    d = np.sqrt(((xx - 384) / 1100) ** 2 + ((yy - 144) / 650) ** 2)
    alpha = np.clip((d - 0.15) / 0.7, 0, 1) * 0.75
    overlay = np.zeros((720, 1280, 4))
    overlay[..., :3] = [14 / 255, 19 / 255, 24 / 255]
    overlay[..., 3] = alpha
    ax.imshow(overlay, extent=(0, 1280, 720, 0), zorder=3)
    if soil:
        top = 600
        ax.add_patch(Rectangle((0, top), 1280, 26, color="#5A3F28", zorder=4))
        ax.add_patch(Rectangle((0, top + 26), 1280, 42, color="#47321F", zorder=4))
        ax.add_patch(Rectangle((0, top + 68), 1280, 60, color="#352517", zorder=4))
        ax.plot([0, 1280], [top, top], color="#8A6440", linewidth=3, zorder=5)
        for cx, cy, r in ((154, top + 48, 4), (474, top + 84, 3), (806, top + 42, 4), (1126, top + 78, 3)):
            ax.add_patch(Circle((cx, cy), r, color="#6E5238", zorder=5))
        for cx, depth in ((918, 6), (1020, 46), (1122, 18)):  # burrows, as on the HTML slides
            ax.add_patch(matplotlib.patches.Ellipse((cx, top + depth), 46, 40, color="#" + EMPTY, zorder=6))
    fig.savefig(path, dpi=100)
    plt.close(fig)


# ---------- drawing helpers ----------

def background(slide, path):
    slide.shapes.add_picture(path, 0, 0, px(1280), px(720))


def box(slide, x, y, w, h, fill=PANEL, line=BORDER, radius=0.06):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(x), px(y), px(w), px(h))
    shape.adjustments[0] = radius
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def text(slide, x, y, w, h, runs, size=20, color=INK, font=BODY_FONT, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.15):
    """runs: a string, or a list of paragraphs, each a string or a list of (text, {bold, color, italic, size}) runs."""
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paragraphs = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        for content, style in (para if isinstance(para, list) else [(para, {})]):
            r = p.add_run()
            r.text = content
            r.font.name = style.get("font", font)
            r.font.size = Pt(style.get("size", size))
            r.font.bold = style.get("bold", bold)
            r.font.italic = style.get("italic", False)
            r.font.color.rgb = rgb(style.get("color", color))
    return tb


def bullets(slide, x, y, w, h, items, size=18, color=INK, gap=6, bullet_color=MUTED, indent=18):
    """items: list of paragraphs (as for text()), each drawn with a bullet; a (level, paragraph) tuple indents it."""
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        level, para = item if isinstance(item, tuple) else (0, item)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.12
        p.space_after = Pt(gap)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(px(indent + level * 22))))
        pPr.set("indent", str(-int(px(indent))))
        clr = pPr.makeelement(qn("a:buClr"), {})
        srgb = clr.makeelement(qn("a:srgbClr"), {"val": bullet_color})
        clr.append(srgb)
        pPr.append(clr)
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•" if level == 0 else "◦"}))
        for content, style in (para if isinstance(para, list) else [(para, {})]):
            r = p.add_run()
            r.text = content
            r.font.name = BODY_FONT
            r.font.size = Pt(style.get("size", size - 2 * level))
            r.font.bold = style.get("bold", False)
            r.font.italic = style.get("italic", False)
            r.font.color.rgb = rgb(style.get("color", color))
    return tb


def kicker(slide, label, color=PLANT, y=50):
    text(slide, 64, y, 1100, 22, label.upper(), size=12, color=color, font=PIXEL_FONT, bold=True)


def title(slide, words, y=74, size=32, w=1150, h=60):
    text(slide, 64, y, w, h, words, size=size, bold=True)


def footer(slide, note, page, color=FOOT):
    text(slide, 64, 692, 1000, 18, note, size=10, color=color)
    text(slide, 1100, 692, 116, 18, "%02d" % page, size=10, color=color, font=PIXEL_FONT, bold=True, align=PP_ALIGN.RIGHT)


def tag(slide, x, y, label, color, w=None):
    w = w or (len(label) * 9.6 + 22)
    shape = box(slide, x, y, w, 26, fill=EMPTY, line=color, radius=0.15)
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label.upper()
    r.font.name = PIXEL_FONT
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = rgb(color)
    return w


def picture(slide, name, x, y, w=None, h=None):
    path = os.path.join(FIG, name)
    iw, ih = Image.open(path).size
    if w and not h:
        h = w * ih / iw
    elif h and not w:
        w = h * iw / ih
    elif w and h:  # fit inside the box, centred
        scale = min(w / iw, h / ih)
        nw, nh = iw * scale, ih * scale
        x, y, w, h = x + (w - nw) / 2, y + (h - nh) / 2, nw, nh
    slide.shapes.add_picture(path, px(x), px(y), px(w), px(h))
    return x, y, w, h


def organism(slide, kind, cx, top, cell=22, gap=3):
    """Draws an organism cell by cell, centred on cx, starting at top. Returns its height."""
    cells = BODIES[kind]
    cols = [c for c, _, _ in cells]
    rows = [r for _, r, _ in cells]
    c0, r0 = min(cols), min(rows)
    width = (max(cols) - c0 + 1) * (cell + gap) - gap
    left = cx - width / 2
    for c, r, kind_ in cells:
        x = left + (c - c0) * (cell + gap)
        y = top + (r - r0) * (cell + gap)
        sq = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(x), px(y), px(cell), px(cell))
        sq.adjustments[0] = 0.14
        sq.fill.solid()
        sq.fill.fore_color.rgb = rgb(CELL[kind_])
        sq.line.fill.background()
        sq.shadow.inherit = False
        if kind_ == "eye":
            slit = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(x + cell * 0.42), px(y + cell * 0.15), px(cell * 0.16), px(cell * 0.7))
            slit.fill.solid()
            slit.fill.fore_color.rgb = rgb(EMPTY)
            slit.line.fill.background()
            slit.shadow.inherit = False
    return (max(rows) - r0 + 1) * (cell + gap) - gap


def specimen(slide, kind, cx, top, label, color, note=None, cell=22, body_h=None):
    h = organism(slide, kind, cx, top + ((body_h - _height(kind, cell)) if body_h else 0), cell)
    y = top + (body_h or h) + 10
    text(slide, cx - 120, y, 240, 20, label.upper(), size=11, color=color, font=PIXEL_FONT, bold=True, align=PP_ALIGN.CENTER)
    if note:
        text(slide, cx - 120, y + 22, 240, 20, note, size=12, color=MUTED, align=PP_ALIGN.CENTER)


def _height(kind, cell, gap=3):
    rows = [r for _, r, _ in BODIES[kind]]
    return (max(rows) - min(rows) + 1) * (cell + gap) - gap


def stat_card(slide, x, y, w, h, big, color, body):
    box(slide, x, y, w, h)
    text(slide, x + 20, y + 14, w - 40, 44, big, size=30, color=color, font=PIXEL_FONT, bold=True)
    text(slide, x + 20, y + 62, w - 40, h - 70, body, size=15)


def B(t, **kw):
    return (t, dict(bold=True, **kw))


def N(t, **kw):
    return (t, dict(kw))


# ---------- slides ----------

def build():
    bg = os.path.join(FIG, "_bg.png")
    bg_soil = os.path.join(FIG, "_bg_soil.png")
    make_background(bg)
    make_background(bg_soil, soil=True)

    prs = Presentation()
    prs.slide_width, prs.slide_height = px(1280), px(720)
    blank = prs.slide_layouts[6]
    page = [0]

    def new(soil=False):
        s = prs.slides.add_slide(blank)
        background(s, bg_soil if soil else bg)
        page[0] += 1
        return s

    # 1. Title
    s = new(soil=True)
    kicker(s, "UCT Honours · Artificial Life · Life Engine", y=96)
    text(s, 64, 128, 760, 240, ["DIGGING", "THEIR OWN", "GRAVES"], size=58, bold=True, font=PIXEL_FONT, spacing=0.95)
    text(s, 64, 410, 740, 80, "Burrowing, soil and extinction: can a trait that saves each individual doom the population?",
         size=22, color=SOFT)
    for kind, cx, label, color in (("plant", 862, "plant", PLANT), ("clawprey", 1008, "burrower", CLAW), ("predator", 1172, "predator", KILLER)):
        specimen(s, kind, cx, 150, label, color, cell=30, body_h=102)
    footer(s, "Group final presentation", page[0], color="A88A6A")

    # 2. The world
    s = new()
    kicker(s, "The organism")
    title(s, "A world of cells, soil and one new trait")
    bullets(s, 64, 168, 660, 430, [
        [B("Organisms are built from cells. "), N("Each cell type does one job, and offspring mutate by adding, removing or changing a cell.")],
        [B("Soil fertility. ", color=SOIL), N("Plants are only born where the soil under them averages at least 7 of 10. Unoccupied soil slowly regains fertility.")],
        [B("The claw cell. ", color=CLAW), N("A prey with a claw that sees a predator "), B("burrows"), N(": it is invisible for 100 ticks and doesn't age, but each burrow strips 10 fertility from the soil under it.")],
        [B("Roles come from the body. "), N("A killer cell and no producer makes a "), N("predator", color=KILLER), N("; a producer that can't move is a "), N("plant", color=PLANT), N("; everything else is "), N("prey", color=PREY), N(".")],
    ], size=19, gap=10)
    lx = 64
    for name in ("producer", "mouth", "mover", "eye", "killer", "claw", "food"):
        sq = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(lx), px(626), px(14), px(14))
        sq.fill.solid()
        sq.fill.fore_color.rgb = rgb(CELL[name])
        sq.line.fill.background()
        text(s, lx + 20, 623, 90, 20, name, size=13, color=MUTED)
        lx += 20 + len(name) * 8 + 26
    box(s, 770, 220, 446, 340)
    specimen(s, "plant", 885, 255, "plant", PLANT, "makes food", body_h=72)
    specimen(s, "prey", 1100, 255, "prey", PREY, "eats food, flees", body_h=72)
    specimen(s, "clawprey", 885, 410, "clawed prey", CLAW, "can burrow", body_h=72)
    specimen(s, "predator", 1100, 410, "predator", KILLER, "eats prey only", body_h=72)
    footer(s, "Starting organisms as built in the simulator (cell layouts from RandomOrganismGenerator)", page[0])

    # 3. The biology
    s = new()
    kicker(s, "The biology")
    title(s, [[N("Evolutionary suicide: when adaptation "), N("causes", color=CLAW), N(" extinction")]])
    text(s, 64, 160, 580, 170, [[N("Natural selection favours whatever helps individuals leave more offspring, "), B("even when it harms the population"),
                                 N(". If a favoured trait degrades something everyone depends on, selection can drive its own carriers extinct.")]], size=19)
    text(s, 64, 340, 580, 70, [[N("Gyllenberg & Parvinen (2001), "), N("Bull. Math. Biol.", italic=True), N("; Rankin & López-Sepulcre (2005), "),
                                N("Oikos", italic=True), N(", \"Can adaptation lead to extinction?\"; Hardin (1968), \"The tragedy of the commons\", "),
                                N("Science", italic=True), N(".")]], size=13, color=MUTED)
    box(s, 64, 430, 580, 190)
    text(s, 86, 450, 540, 30, "A real echo: burrowers reshape soil", size=19, bold=True, color=SOIL)
    text(s, 86, 488, 540, 120, "Burrowing animals are ecosystem engineers. Introduced European rabbits in Australia, for example, burrow and graze so heavily that they strip vegetation and drive soil erosion.", size=15)
    box(s, 684, 160, 532, 460)
    text(s, 706, 180, 480, 30, "Our hypothesis", size=21, bold=True)
    steps = [("1 · individual gain", "Burrowing hides a prey from predators"), ("2 · selection", "Claws spread through the prey"),
             ("3 · shared cost", "Every burrow drains the soil plants need"), ("4 · collapse", "Plants fail, food runs out, burrowers die")]
    for i, (head, body) in enumerate(steps):
        y = 226 + i * 78
        box(s, 706, y, 488, 70, fill=EMPTY, radius=0.1)
        text(s, 722, y + 10, 460, 18, head.upper(), size=11, color=CLAW, font=PIXEL_FONT, bold=True)
        text(s, 722, y + 34, 460, 26, body, size=16)
    text(s, 706, 548, 488, 60, [[N("The test: does the population survive "), B("without"), N(" the trait, and only collapse "), B("with"), N(" it?")]], size=15, color=MUTED)
    footer(s, "Mis-adaptation = a trait that is favoured yet ruinous · failure to adapt = no suitable trait in time", page[0])

    # 4. The experiment
    s = new()
    kicker(s, "The experiment")
    title(s, "Nearly 2,000 simulations, run like a field study")
    bullets(s, 64, 190, 660, 420, [
        [B("Same code, no screen. "), N("A headless runner drives the simulator's own code, so hundreds of runs go in parallel.")],
        [B("Seeded layouts, random lives. "), N("Soil and starting positions come from seeds; everything after that is stochastic, so replicates differ.")],
        [B("Controls, like knockout experiments: "), N("the same starting layouts with one thing changed.")],
        (1, [N("claws "), B("seeded"), N(" at the start, or left to "), B("evolve"), N(" by mutation")]),
        (1, [N("burrowing "), B("switched off"), N(", or "), B("harmless"), N(" to the soil")]),
        (1, [N("soil density, patch size, regrowth; predator strength; plant regrowth")]),
    ], size=19, gap=10)
    box(s, 770, 196, 446, 400)
    rows = [("Study", "Focus", "Runs"), ("Tarique", "evolved claws, soil cost, predators", "456"), ("Nik", "evolved claws, soil settings, maps", "170+"),
            ("Jenna", "plant regrowth, plant death", "936"), ("Josh", "seeded vs evolved claws, soil layouts", "380")]
    for i, (a, b, c) in enumerate(rows):
        y = 214 + (0 if i == 0 else 36 + (i - 1) * 64)
        style = dict(size=13 if i == 0 else 15, color=MUTED if i == 0 else INK, bold=i == 0)
        text(s, 790, y, 84, 56, a, **style)
        text(s, 874, y, 250, 56, b, **style)
        text(s, 1130, y, 66, 56, c, align=PP_ALIGN.RIGHT, **style)
        if i < len(rows) - 1:
            line_y = y + (28 if i == 0 else 56)
            ln = s.shapes.add_connector(1, px(790), px(line_y), px(1196), px(line_y))
            ln.line.color.rgb = rgb(BORDER)
    text(s, 790, 540, 410, 40, "Plus browser runs. 100,000–250,000 ticks each. Proportions compared with Fisher's exact test.", size=12, color=MUTED)
    footer(s, "Settings differ slightly between studies; each comparison is made within one study", page[0])

    # 5. Q1
    s = new()
    kicker(s, "Question 1 · Did an extinction event occur?", color=CLAW)
    title(s, "Yes, at every level, and often")
    stat_card(s, 64, 160, 476, 150, "167/380", KILLER, [[N("of our headless runs ended with "), B("every organism dead"), N("; in Nik's sweep, 111 of 111")]])

    stat_card(s, 64, 326, 476, 150, "21/40", CLAW, [[N("runs where evolved claws took over and then "), B("every prey died"), N(" (Tarique)")]])
    stat_card(s, 64, 492, 476, 170, "LINEAGES", PREY, [[N("The seeded predators died out in "), B("every"), N(" batch run; seeded burrowing prey died out in every run with all the claw's benefits switched off")]])
    _, y, _, h = picture(s, "extinction_timing.png", 576, 200, w=640)
    text(s, 576, y + h + 10, 640, 40, "When 160 worlds died (batches with seeded claws, 300 runs). Most die early, a few after 100k ticks of apparent stability.", size=12, color=MUTED)
    footer(s, "Whole-world extinction usually ends with plants: with no animals to eat it, food piles up and blocks plant offspring", page[0])

    # 6. Q2 failure to adapt
    s = new()
    kicker(s, "Question 2 · Mis-adaptation or failure to adapt?", color=CLAW)
    w = tag(s, 64, 92, "failure to adapt", KILLER)
    title(s, "Most worlds die before evolution can act", y=82, w=1100 - w, size=30)
    s.shapes[-1].left = px(64 + w + 14)
    bullets(s, 64, 190, 470, 440, [
        [B("The first predator wave. "), N("Seeded predators feed on the first prey boom, multiply about 25×, and overshoot.")],
        [B("88 of 160"), N(" world extinctions came before 30,000 ticks. When the prey died, "), B("82%"), N(" of predators were still the seeded ones.")],
        [N("Claws appear within ~1,200 ticks, but prey were usually gone by "), B("5,500"), N(", while claws were on only ~4% of them (Nik).")],
        [N("With all the claw's benefits off, the seeded burrowers bred at half the rate of other prey and were the first to vanish.")],
    ], size=18, gap=10)
    _, y, _, h = picture(s, "early_wave.png", 566, 230, w=650)
    text(s, 566, y + h + 10, 650, 30, [[N("Browser run, claw benefits off: "), N("predators", color=KILLER), N(" boom and the "), N("clawed prey", color=CLAW), N(" are gone by tick 4,000.")]], size=12, color=MUTED)
    footer(s, "Ecological overshoot, not evolution: the prey had no time to adapt", page[0])

    # 7. Q2 claw evolves
    s = new()
    kicker(s, "Question 2 · Mis-adaptation or failure to adapt?", color=CLAW)
    w = tag(s, 64, 92, "mis-adaptation", CLAW)
    title(s, "Given time, the claw evolves, because it works", y=82, w=1100 - w, size=30)
    s.shapes[-1].left = px(64 + w + 14)
    picture(s, "claw_spread.png", 64, 150, w=1152, h=400)
    text(s, 64, 568, 560, 90, [[B("No organism starts with a claw. "), N("It appears by mutation in "), B("every run"), N(" (80 of 80), within a median of about 1,000 ticks.")]], size=16)
    text(s, 660, 568, 556, 90, [[B("It spreads only when burrowing works: "), N("7/10 take-overs against 1/10 when claws cost food (p = 0.02), 10/10 against 2/10 when free (p = 0.0007). Selection favours the claw "), N("for", italic=True), N(" burrowing.")]], size=16)
    footer(s, "Josh's claw-evolution batch · 10 runs per setup · predator lifespan multiplier 2 · bands are ±1 std", page[0])

    # 8. Q2 collapse
    s = new()
    kicker(s, "Question 2 · Mis-adaptation or failure to adapt?", color=CLAW)
    w = tag(s, 64, 92, "mis-adaptation", CLAW)
    title(s, "…and then the burrowers dig their own grave", y=82, w=1100 - w, size=30)
    s.shapes[-1].left = px(64 + w + 14)
    chain = [("8k", "Evolved claws take over the prey"), ("8k → 15k", "Burrows drain the soil to 74%"), ("→ 20k", "Plants crash from 468 to 35"),
             ("→ 57k", "Clawed prey dwindle and die"), ("65k", "Nothing is left")]
    cw = 1152 / 5
    for i, (head, body) in enumerate(chain):
        x = 64 + i * cw
        box(s, x, 150, cw - 6, 84, radius=0.08)
        text(s, x + 12, 160, cw - 30, 18, head.upper(), size=11, color=CLAW, font=PIXEL_FONT, bold=True)
        text(s, x + 12, 182, cw - 30, 48, body, size=14)
    stat_card(s, 64, 254, 476, 176, "21/40 vs 0/20", CLAW, [[N("prey extinct after claws took over, damaging vs "), B("harmless"), N(" burrowing (Tarique, p = 0.00002). Without working claws, prey survived 17/20.")]])
    stat_card(s, 64, 446, 476, 160, "11/11 vs 0/5", CLAW, "burrower populations extinct after spreading, with vs without soil damage (Nik, p ≈ 0.0002).")
    picture(s, "collapse_chain.png", 572, 254, w=644, h=360)
    footer(s, "One run (claws evolved, free, burrowing on). The soil recovers after the burrowers die, too late for them.", page[0])

    # 9. Q2 verdict
    s = new()
    kicker(s, "Question 2 · Mis-adaptation or failure to adapt?", color=CLAW)
    title(s, "Both, depending on how far evolution gets")
    for x, label, color, head, items in (
            (64, "failure to adapt", KILLER, "The common route",
             ["Strong seeded predators, prey with no defence yet", "Collapse within ~30k ticks, before the claw spreads",
              "Happens just as often with burrowing switched off (Nik's controls)"]),
            (660, "mis-adaptation", CLAW, "Evolutionary suicide, when the claw gets a chance",
             [[N("The claw evolves and spreads "), B("because"), N(" it protects (vs burrowing off)")],
              "It kills only when burrowing damages the soil (vs harmless burrowing)",
              "Prey survive without the claw: the collapse is the trait's doing"])):
        box(s, x, 154, 556, 318)
        tag(s, x + 22, 174, label, color)
        text(s, x + 22, 212, 512, 60, head, size=19, bold=True)
        bullets(s, x + 22, 276 if len(head) > 30 else 250, 512, 190, items, size=15, gap=5)
    box(s, 64, 488, 1152, 172)
    text(s, 86, 504, 1100, 30, "A second, subtler route", size=20, bold=True, color=KILLER)
    text(s, 86, 540, 1108, 110, [[N("Clawed prey that grow a killer cell can become predators. In 9 of 22 late collapses, the predator that wiped out the prey had evolved "),
                                  B("from the prey's own lineage"), N(", and in about 10 a newly evolved predator spread and the world died within a few thousand ticks. Suggestive of evolutionary suicide, but not yet proven by replay.")]], size=16)
    footer(s, "Not yet measured: whether clawed prey still out-reproduce clawless prey as the population declines", page[0])

    # 10. Q3
    s = new()
    kicker(s, "Question 3 · Environment vs the initial seeded species?", color=CLAW)
    title(s, [[N("The seeded species decide "), N("whether", color=CLAW), N(" burrowers get a chance; the environment decides "), N("what happens", color=CLAW), N(" to them")]], size=29, h=90)
    tag(s, 64, 186, "seeded species", PREY)
    bullets(s, 64, 222, 480, 220, [
        [B("Predator strength"), N(" sets the first wave: at lifespan ×3, predators killed all prey in 7/8 runs even without burrowing; ×2 lets prey survive (Tarique)")],
        [B("Seeding claws"), N(" lets prey skip the wait and survive the first wave")],
        [N("The seeded predators "), B("never last"), N(": evolved predators replace them")],
    ], size=16, gap=6)
    tag(s, 64, 470, "environment", PLANT)
    bullets(s, 64, 506, 480, 170, [
        [B("Soil"), N(" density, patch size and regrowth (right)")],
        [B("Plant regrowth"), N(" rescues worlds from the claw (Jenna)")],
        [B("Map luck"), N(": the pattern appeared on one soil map, not another (Nik)")],
    ], size=16, gap=6)
    _, y, _, h = picture(s, "soil_survival.png", 570, 200, w=646)
    text(s, 570, y + h + 10, 646, 40, "Josh's soil-layouts batch: 270 runs, 6 setups × 9 soil layouts × 5 soil seeds. Runs lasting 250,000 ticks.", size=12, color=MUTED)
    footer(s, "Same starting species, different soil: survival ranges from 80% to 0%", page[0])

    # 11. Q4
    s = new()
    kicker(s, "Question 4 · Instrumental features and traits", color=CLAW)
    title(s, "What made the difference")
    picture(s, "group_evidence.png", 64, 140, w=1152, h=320)
    text(s, 64, 462, 1152, 20, "Left: Tarique, soil removed per burrow. Middle: Nik, soil damage on vs off. Right: Jenna, plant regrowth (reference settings).",
         size=12, color=MUTED, align=PP_ALIGN.CENTER)
    tag(s, 64, 496, "environment", PLANT)
    bullets(s, 64, 532, 560, 150, [
        [B("Soil damage per burrow"), N(": 0 or 3 never collapses; 10 collapses half")],
        [B("Soil and plant regrowth"), N(": fast recovery rescues the world")],
        [B("Food blocking"), N(": without consumers, food chokes plant offspring")],
    ], size=15, gap=4)
    tag(s, 660, 496, "traits", CLAW)
    bullets(s, 660, 532, 556, 150, [
        [B("Claw", color=CLAW), N(": the trait that both saves and dooms")],
        [B("Eye"), N(": claw mutations often replace it; 11 of 12 clawed prey in one replay were blind and never burrowed")],
        [B("Killer cell", color=KILLER), N(": turns clawed prey into predators; "), B("hide length"), N(": at 30 ticks the claw never pays")],
    ], size=15, gap=4)
    footer(s, "A window of settings: predators strong enough to make burrowing pay, burrows damaging enough to matter", page[0])

    # 12. Loophole
    s = new()
    kicker(s, "Evolution finds loopholes")
    title(s, [[N("When claws were free, "), N("predators", color=CLAW), N(" took them, and lost their eyes")]], size=30)
    specimen(s, "predator", 170, 200, "seeded predator", KILLER, "6 cells · 11 food to breed", body_h=72)
    text(s, 290, 214, 60, 40, "→", size=30, color=MUTED, align=PP_ALIGN.CENTER)
    specimen(s, "eyeless", 480, 200, "eyeless clawed predator", CLAW, "5 cells · 4 food to breed", body_h=72)
    bullets(s, 64, 350, 640, 300, [
        [N("A "), B("free claw"), N(" adds lifespan and health but costs no food, so predators keep them; none ever burrowed.")],
        [N("Mouths are useless to predators and eyes cost food: a "), B("cheap random walker"), N(" breeding ~3× as often beats a hunter.")],
        [N("It came to dominate some worlds on its own: selection optimises the rules, not the design.")],
    ], size=17, gap=8)
    _, y, _, h = picture(s, "predator_claws.png", 740, 200, w=476)
    text(s, 740, y + h + 10, 476, 40, "Share of predators carrying claws late in the runs, across 9 soil layouts. Claw-off: claws cost food.", size=12, color=MUTED)
    footer(s, "Eyeless predator layout is illustrative; its cells (mover, killer, 3 claws) are from the species record", page[0])

    # 13. Answers
    s = new()
    kicker(s, "Answers to the brief")
    title(s, "In four lines")
    answers = [("Did an extinction event occur?", "Yes: whole worlds, prey populations and lineages went extinct in a large share of runs."),
               ("Mis-adaptation or failure to adapt?", "Mostly failure to adapt, as the first predator wave struck before the claw spread. When the claw had time, it was a mis-adaptation: favoured for hiding, fatal through the soil it destroyed."),
               ("Environment vs the initial seeded species?", "The seeded predators decide whether burrowers get a chance; soil and plant recovery decide whether the burrowers then destroy themselves."),
               ("Instrumental features and traits?", "Soil damage per burrow, soil and plant regrowth, food blocking; the claw, the eye it replaces, and the killer cell that turns prey into predators.")]
    for i, (q, a) in enumerate(answers):
        y = 160 + i * 122
        text(s, 64, y, 50, 40, str(i + 1), size=28, color=CLAW, font=PIXEL_FONT, bold=True)
        text(s, 130, y, 1086, 30, q, size=20, bold=True)
        text(s, 130, y + 34, 1086, 70, a, size=17, color=SOFT)
        if i < 3:
            ln = s.shapes.add_connector(1, px(64), px(y + 110), px(1216), px(y + 110))
            ln.line.color.rgb = rgb(BORDER)
    footer(s, "", page[0])

    # 14. Limitations
    s = new(soil=True)
    kicker(s, "Limitations and next steps")
    title(s, "What we can't claim yet")
    for x, head, items in (
            (64, "Limitations", ["10–40 runs per setup: differences of 2–3 runs could be chance", "Studies used slightly different settings and grids",
                                 "Collapse rate depends on behaviour settings: with \"burrow only when a predator is close\" on, evolved claws collapsed 2 of 20 worlds here, against 21 of 40 in Tarique's runs with it off",
                                 "Proximate cause of death is mostly predation; the soil acts by stopping plants recovering"]),
            (660, "Next", [[N("Measure fitness of clawed vs clawless prey "), B("during"), N(" a decline: does selection still favour the claw?")],
                           "Replay suspected suicides with and without the invading predator", "Close the free-claw loophole for predators", "More soil maps per setup"])):
        box(s, x, 170, 556, 380)
        text(s, x + 22, 190, 512, 30, head, size=21, bold=True)
        bullets(s, x + 22, 232, 512, 310, items, size=16, gap=8)
    footer(s, "", page[0], color="A88A6A")

    # 15. Thanks
    s = new(soil=True)
    kicker(s, "Thank you", y=118)
    text(s, 64, 146, 700, 80, "QUESTIONS?", size=48, bold=True, font=PIXEL_FONT)
    text(s, 64, 236, 820, 100, "Built on The Life Engine by Max Robinson, extended with soil fertility, burrowing claws, role tagging, seeded run setup and a headless experiment runner.", size=19, color=SOFT)
    text(s, 64, 356, 860, 110, [[N("References: Gyllenberg, M. & Parvinen, K. (2001). Necessary and sufficient conditions for evolutionary suicide. "),
                                 N("Bulletin of Mathematical Biology", italic=True), N(" 63, 981–993. · Rankin, D. J. & López-Sepulcre, A. (2005). Can adaptation lead to extinction? "),
                                 N("Oikos", italic=True), N(" 111, 616–619. · Hardin, G. (1968). The tragedy of the commons. "), N("Science", italic=True), N(" 162, 1243–1248.")]],
         size=14, color=MUTED)
    organism(s, "clawprey", 1046, 124, cell=28)
    organism(s, "plant", 1180, 124, cell=28)
    footer(s, "Reports: Burrowing Into Extinction (Tarique) · Soil study (Nik) · Plant regrowth (Jenna) · Claw evolution and soil layouts (Josh)", page[0], color="A88A6A")

    prs.save(OUT)
    print("wrote", OUT, "with", len(prs.slides), "slides")


if __name__ == "__main__":
    build()
